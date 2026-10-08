"""Persistencia de noticias en Excel."""

from __future__ import annotations

from datetime import datetime, date
from pathlib import Path
from collections.abc import Iterable

from openpyxl import Workbook, load_workbook

from news_scraper.models.noticia import Noticia
from news_scraper.services.validator import remove_duplicates
from news_scraper.config import DATE_FORMAT


HEADERS = ("dia", "fuente", "titulo", "parrafo")


class ExcelService:
    """Crea y actualiza el archivo Excel de noticias."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.last_duplicates = 0

    def _ensure_workbook(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(HEADERS)
            workbook.save(self.path)
        return load_workbook(self.path)

    def existing_keys(self) -> set[tuple[date, str, str]]:
        """Lee las claves de noticias ya guardadas."""

        workbook = self._ensure_workbook()
        sheet = workbook.active
        keys: set[tuple[date, str, str]] = set()
        for row in sheet.iter_rows(min_row=2, max_col=3):
            if not all(cell.value is not None for cell in row):
                continue
            raw_date = row[0].value
            if isinstance(raw_date, datetime):
                parsed_date = raw_date.date()
            elif isinstance(raw_date, date):
                parsed_date = raw_date
            else:
                try:
                    parsed_date = datetime.strptime(str(raw_date), DATE_FORMAT).date()
                except ValueError:
                    continue
            keys.add((parsed_date, str(row[1].value), str(row[2].value)))
        workbook.close()
        return keys

    def append_news(self, news_items: Iterable[Noticia]) -> int:
        """Agrega noticias nuevas y devuelve cuántas filas se añadieron."""

        items = list(news_items)
        unique_items = remove_duplicates(items)
        self.last_duplicates = len(items) - len(unique_items)
        workbook = self._ensure_workbook()
        sheet = workbook.active
        existing = self.existing_keys()
        added = 0

        for news in unique_items:
            if news.duplicate_key() in existing:
                self.last_duplicates += 1
                continue
            sheet.append(
                (
                    news.dia.strftime(DATE_FORMAT),
                    news.fuente,
                    news.titulo,
                    news.parrafo,
                )
            )
            existing.add(news.duplicate_key())
            added += 1

        workbook.save(self.path)
        workbook.close()
        return added

