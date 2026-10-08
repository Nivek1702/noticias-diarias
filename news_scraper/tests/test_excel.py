from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from news_scraper.models.noticia import Noticia
from news_scraper.services.excel import ExcelService, HEADERS


def test_excel_crea_columnas_y_evita_duplicados() -> None:
    path = Path(__file__).with_name("_noticias_test.xlsx")
    service = ExcelService(path)
    noticia = Noticia(date(2026, 10, 8), "Portafolio", "Título", "Párrafo")

    try:
        assert service.append_news([noticia, noticia]) == 1
        assert service.append_news([noticia]) == 0

        workbook = load_workbook(path)
        sheet = workbook.active
        assert tuple(sheet.values)[0] == HEADERS
        assert sheet.max_row == 2
        assert sheet.cell(row=2, column=1).value == "08/10/2026"
        workbook.close()
    finally:
        path.unlink(missing_ok=True)

